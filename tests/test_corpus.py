"""P3 corpus readiness: acquisition, batches, file routing, intake, readiness (synthetic engineering fixtures)."""

from __future__ import annotations

import hashlib
import io
import json
import zipfile
from pathlib import Path
from types import SimpleNamespace

import httpx
import pandas as pd
import pytest
import yaml

from regkg.config import BudgetAssumptions, ReadyBatchLimits
from regkg.literature import files
from regkg.literature.acquire import acquire_batch
from regkg.literature.advisors import AdvisorReconciliationError, reconcile_advisors
from regkg.literature.http import NcbiSettings, SourceClient
from regkg.literature.intake import run_intake, safe_url
from regkg.literature.parse import Document, parse_bioc
from regkg.workflows import corpus as c
from regkg.workflows import corpus_ready as r

S3 = "http://s3.amazonaws.com/doc/2006-03-01/"
JATS = (
    b"<article><front><article-meta><title-group><article-title>TEAD1 in lung</article-title></title-group>"
    b"</article-meta></front><body><sec><title>Results</title><p>TEAD1 binds.</p>"
    b"<supplementary-material id='s1'><label>Table S1</label><media xlink:href='supp1.xlsx' "
    b"xmlns:xlink='http://www.w3.org/1999/xlink'/></supplementary-material>"
    b"<supplementary-material id='s2'><label>Table S2</label><media xlink:href='missing.xlsx' "
    b"xmlns:xlink='http://www.w3.org/1999/xlink'/></supplementary-material></sec></body></article>"
)


def pubmed_xml(pmids: list[str], pmc: dict[str, str]) -> bytes:
    items = "".join(
        f"<PubmedArticle><MedlineCitation><PMID>{p}</PMID><Article><ArticleTitle>Paper {p}</ArticleTitle>"
        f"<Abstract><AbstractText>TEAD1 knockdown reduced target promoter activity in lung cells.</AbstractText>"
        f"</Abstract></Article></MedlineCitation><PubmedData><ArticleIdList>"
        + (f"<ArticleId IdType='pmc'>{pmc[p]}</ArticleId>" if p in pmc else "")
        + f"<ArticleId IdType='doi'>10.1/p{p}</ArticleId></ArticleIdList></PubmedData></PubmedArticle>"
        for p in pmids
    )
    return f"<PubmedArticleSet>{items}</PubmedArticleSet>".encode()


def s3_listing(prefixes=(), files_=()) -> bytes:
    body = "".join(f"<CommonPrefixes><Prefix>{p}</Prefix></CommonPrefixes>" for p in prefixes)
    body += "".join(f"<Contents><Key>{k}</Key><Size>{s}</Size></Contents>" for k, s in files_)
    return f"<ListBucketResult xmlns='{S3}'><IsTruncated>false</IsTruncated>{body}</ListBucketResult>".encode()


class Sources:
    """Mock PubMed, Europe PMC, PubTator, and PMC OA dataset for three papers:
    1 = OA (Europe PMC JATS fails with HTTP 500; recovered from the OA dataset), 2 = free-to-read not OA,
    3 = PubMed timeout for the batch is simulated separately."""

    def __init__(self, supplement=b"PK-not-really", supplement_md5=None, epmc_status=500):
        self.supplement, self.epmc_status = supplement, epmc_status
        self.supplement_md5 = supplement_md5 or hashlib.md5(supplement).hexdigest()
        self.jats_md5 = hashlib.md5(JATS).hexdigest()
        self.requests: list[str] = []

    def __call__(self, request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        self.requests.append(url)
        if "efetch" in url:
            return httpx.Response(200, content=pubmed_xml(["1", "2"], {"1": "PMC1", "2": "PMC2"}))
        if "fullTextXML" in url:
            return httpx.Response(self.epmc_status, content=JATS if self.epmc_status == 200 else b"")
        if "europepmc" in url and "search" in url:
            result = [
                {
                    "pmid": "1",
                    "pmcid": "PMC1",
                    "isOpenAccess": "Y",
                    "inPMC": "Y",
                    "hasSuppl": "Y",
                    "source": "MED",
                    "id": "1",
                    "fullTextUrlList": {
                        "fullTextUrl": [
                            {
                                "availability": "Open access",
                                "documentStyle": "pdf",
                                "site": "Europe_PMC",
                                "url": "https://europepmc.org/a1",
                            }
                        ]
                    },
                },
                {
                    "pmid": "2",
                    "pmcid": "PMC2",
                    "isOpenAccess": "N",
                    "inPMC": "Y",
                    "source": "MED",
                    "id": "2",
                    "fullTextUrlList": {
                        "fullTextUrl": [
                            {
                                "availability": "Free",
                                "documentStyle": "pdf",
                                "site": "Europe_PMC",
                                "url": "https://europepmc.org/a2?pdf=render",
                            }
                        ]
                    },
                },
            ]
            return httpx.Response(200, json={"hitCount": 2, "resultList": {"result": result}})
        if "pubtator" in url:
            return httpx.Response(200, json={"PubTator3": []})
        if "pmc-oa-opendata" in url:
            path = request.url.path
            prefix = request.url.params.get("prefix")
            if prefix == "PMC1.":
                return httpx.Response(200, content=s3_listing(prefixes=["PMC1.1/"]))
            if prefix == "PMC1.1/":
                return httpx.Response(
                    200,
                    content=s3_listing(
                        files_=[
                            ("PMC1.1/PMC1.1.xml", len(JATS)),
                            ("PMC1.1/supp1.xlsx", len(self.supplement)),
                            ("PMC1.1/fig1.jpg", 10),
                        ]
                    ),
                )
            if prefix:
                return httpx.Response(200, content=s3_listing())
            if path.endswith("PMC1.1.json"):
                return httpx.Response(
                    200,
                    json={
                        "pmcid": "PMC1",
                        "pmid": 1,
                        "doi": "10.1/p1",
                        "is_pmc_openaccess": True,
                        "is_retracted": False,
                        "license_code": "CC BY",
                        "xml_url": f"s3://pmc-oa-opendata/PMC1.1/PMC1.1.xml?md5={self.jats_md5}",
                        "media_urls": [f"s3://pmc-oa-opendata/PMC1.1/supp1.xlsx?md5={self.supplement_md5}"],
                    },
                )
            if path.endswith("PMC1.1.xml"):
                return httpx.Response(200, content=JATS)
            if path.endswith("supp1.xlsx"):
                return httpx.Response(200, content=self.supplement)
        return httpx.Response(404)


def client(tmp_path: Path, handler) -> SourceClient:
    return SourceClient(
        tmp_path / "cache",
        NcbiSettings("tool"),
        timeout_seconds=1,
        max_retries=0,
        transport=httpx.MockTransport(handler),
    )


@pytest.fixture(autouse=True)
def no_sleep(monkeypatch):
    monkeypatch.setattr("time.sleep", lambda _: None)


def acquire(tmp_path, sources, fetch_supplements=True):
    papers = acquire_batch(
        client(tmp_path, sources),
        [("pub:1", "1"), ("pub:2", "2")],
        tmp_path / "papers",
        tmp_path,
        30_000_000,
        ["xlsx", "pdf"],
        fetch_supplements,
        set(),
        {"1": "10.1/p1"},
    )
    return {p.pmid: p.to_json() for p in papers}


def test_open_access_fallback_recovers_and_non_oa_needs_a_manual_download(tmp_path):
    audits = acquire(tmp_path, Sources())
    oa, free = audits["1"], audits["2"]
    assert oa["full_text_route"] == "pmc_oa_dataset_jats"  # Europe PMC HTTP 500 recovered by the dataset copy
    assert c.acquisition_state(oa) == (c.FULL_TEXT_ACQUIRED, "pmc_oa_dataset_jats")
    failures = c.acquisition_failures({"1": oa})
    assert failures["recovered_by"].tolist() == ["pmc_oa_dataset_jats"]  # a system error, not "paywalled"
    supplements = {s["filename"]: s["state"] for s in oa["supplements"]}
    assert supplements == {"supp1.xlsx": "acquired", "missing.xlsx": "not_in_oa_dataset"}  # figures excluded
    state, reason = c.acquisition_state(free)
    assert state == c.MANUAL_DOWNLOAD_NEEDED and "open-access" in reason
    assert c.access_status(free) == "free_to_read_not_open_access"


def test_true_unavailability_differs_from_network_failure(tmp_path):
    def down(request):
        if "efetch" in str(request.url):
            raise httpx.ReadTimeout("slow", request=request)
        return Sources()(request)

    audits = {
        p.pmid: p.to_json()
        for p in acquire_batch(client(tmp_path, down), [("pub:1", "1")], tmp_path / "p", tmp_path, 1, [], False, set())
    }
    assert audits["1"]["pubmed_status"] == "failed:timeout"
    assert c.acquisition_state(audits["1"])[0] == c.FULL_TEXT_ACQUIRED  # the OA route still worked
    bare = dict(audits["1"], full_text_route=None, full_text_status="not_checked")
    assert c.acquisition_state(bare)[0] == c.FETCH_FAILED  # never MANUAL_DOWNLOAD_NEEDED or "unavailable"


def test_supplement_checksum_mismatch_is_not_acquired(tmp_path):
    audits = acquire(tmp_path, Sources(supplement_md5="0" * 32))
    state = {s["filename"]: s["state"] for s in audits["1"]["supplements"]}["supp1.xlsx"]
    assert state == "md5_mismatch"


def test_resumable_batches_never_repeat_successful_work(tmp_path):
    calls = []

    def process(chunk):
        calls.append(list(chunk))
        if chunk == [3, 4] and len(calls) == 2:
            raise RuntimeError("interrupted")
        return [x * 10 for x in chunk], {}

    with pytest.raises(RuntimeError):
        c.run_batches(tmp_path, [1, 2, 3, 4, 5], 2, {"route": "t"}, process)
    assert (tmp_path / "batch-001.json").is_file() and not (tmp_path / "batch-002.json").exists()
    limited = c.run_batches(tmp_path, [1, 2, 3, 4, 5], 2, {"route": "t"}, process, limit=1)
    assert limited["pending_batches"] == 1 and calls[-1] == [3, 4]
    done = c.run_batches(tmp_path, [1, 2, 3, 4, 5], 2, {"route": "t"}, process)
    assert calls == [[1, 2], [3, 4], [3, 4], [5]] and done["pending_batches"] == 0
    assert [x for b in done["records"] for x in b["results"]] == [10, 20, 30, 40, 50]


def test_advisor_reconciliation_fails_on_missing_or_extra_ids(tmp_path):
    source = tmp_path / "list.csv"
    source.write_text("pmid,title\n1,A\n2,B\n")
    digest = hashlib.sha256(source.read_bytes()).hexdigest()

    def seeds(pmids):
        return pd.DataFrame(
            {
                "pmid": pmids,
                "memberships": [[{"source_path": "list.csv", "source_row_index": i}] for i, _ in enumerate(pmids)],
            }
        )

    result = reconcile_advisors(seeds(["1", "2"]), {"list.csv": digest}, tmp_path)
    assert (result.unique_pmids, result.memberships) == (2, 2)
    with pytest.raises(AdvisorReconciliationError, match="pmids_missing_from_p2"):
        reconcile_advisors(seeds(["1"]), {"list.csv": digest}, tmp_path)
    with pytest.raises(AdvisorReconciliationError, match="SHA-256"):
        reconcile_advisors(seeds(["1", "2"]), {"list.csv": "0" * 64}, tmp_path)


def test_generic_bioc_footnote_is_not_a_table_footnote():
    bioc = {
        "passages": [
            {"offset": 0, "text": "Paper", "infons": {"section_type": "TITLE", "type": "front"}},
            {"offset": 10, "text": "Body text here.", "infons": {"section_type": "RESULTS", "type": "paragraph"}},
            {
                "offset": 30,
                "text": "URLs: https://example.org",
                "infons": {"section_type": "SUPPL", "type": "footnote"},
            },
            {
                "offset": 60,
                "text": "Values are means.",
                "infons": {"section_type": "TABLE", "type": "footnote", "id": "T1"},
            },
            {
                "offset": 80,
                "text": "*P < 0.05",
                "infons": {"section_type": "TABLE", "type": "table_footnote", "id": "T1"},
            },
        ]
    }
    kinds = [p.kind for p in parse_bioc("p:x", json.dumps(bioc).encode()).passages]
    assert kinds == ["title", "paragraph", "footnote", "table_footnote", "table_footnote"]


def test_content_sniffing_catches_html_named_pdf_and_legacy_workbooks():
    assert files.detect_type(b"<!DOCTYPE html><html>Sign in</html>", "article.pdf") == (
        "html",
        ["extension_pdf_but_content_html"],
    )
    assert files.detect_type(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100, "t.xls")[0] == "ole2_legacy_office"
    assert files.detect_type(b"", "x.pdf") == ("empty", ["empty_file"])


def xlsx_bytes() -> bytes:
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Results"
    sheet.append(["Gene", "log2FC", "p"])
    sheet.append(["TEAD1", 1.5, "NA"])
    sheet.append(["SEPT2", "=B2*2", "<0.01"])  # formula without a cached value
    sheet.merge_cells("A5:B5")
    hidden = book.create_sheet("README")
    hidden.sheet_state = "hidden"
    hidden.append(["Legend: log2FC is treated vs control"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_spreadsheet_keeps_raw_tokens_formulas_hidden_sheets_and_coordinates():
    document, report = files.parse_xlsx("p:x", xlsx_bytes(), "S1.xlsx", 20, 2000, 60)
    sheets = {s["sheet"]: s for s in report.components["sheets"]}
    assert sheets["README"]["state"] == "hidden" and sheets["README"]["rows_inspected"] == 1  # inventoried and read
    assert sheets["Results"]["merged_ranges"] == ["A5:B5"] and sheets["Results"]["formula_cells"] == 1
    rows = [p for p in document.passages if p.kind == "supplement_table_row" and p.item_id == "Results"]
    cells = {cell["header"]: cell for cell in rows[0].cells}
    assert cells["p"]["value"] == "NA" and cells["log2FC"]["value"] == "1.5"  # "NA" stays a literal token
    formula = next(cell for cell in rows[1].cells if cell["formula"])
    assert formula["formula"] == "=B2*2" and formula["cached_value_missing"]  # never evaluated
    assert rows[1].locator == "xlsx:sheet=Results;row=3"
    assert report.interpretation == "table_meaning_requires_review"


def test_delimited_rows_keep_quoted_newlines_na_and_flag_ragged_rows():
    content = b'gene,value,note\nTEAD1,NA,"line one\nline two"\n007,1e-5\n'
    document, report = files.parse_delimited("p:x", content, "t.csv", ",", 100, 10)
    row = document.passages[1]
    assert [cell["value"] for cell in row.cells] == ["TEAD1", "NA", "line one\nline two"]
    assert document.passages[2].cells[0]["value"] == "007"  # leading zeros preserved
    assert any("non_rectangular" in reason for reason in report.reasons) and report.state == files.PARTIAL


def test_zip_inspection_rejects_traversal_and_symlinks():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("data/table.csv", "a,b\n1,2\n")
        archive.writestr("../escape.txt", "x")
        link = zipfile.ZipInfo("link")
        link.external_attr = 0o120777 << 16
        archive.writestr(link, "/etc/passwd")
    members, report = files.inspect_zip(buffer.getvalue(), "s.zip")
    dispositions = {m["name"]: m["disposition"] for m in report.members}
    assert dispositions == {
        "data/table.csv": "extracted_for_routing",
        "../escape.txt": "rejected_path_traversal",
        "link": "rejected_symlink",
    }
    assert [name for name, _ in members] == ["table.csv"] and report.state == files.PARTIAL


def minimal_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    return out + f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()


def test_intake_checks_identity_bytes_and_urls(tmp_path):
    inbox = tmp_path / "inbox"
    (inbox / "1").mkdir(parents=True)
    (inbox / "2").mkdir()
    (inbox / "1" / "article.pdf").write_bytes(minimal_pdf("PMID 1 Paper about TEAD1"))
    (inbox / "2" / "article.pdf").write_bytes(minimal_pdf("An unrelated document"))
    (inbox / "2" / "fake.pdf").write_bytes(b"<!DOCTYPE html><html>login</html>")
    (inbox / "intake.csv").write_text(
        "pmid,filename,role,source_url,acquired_on,disposition,note\n"
        "1,article.pdf,article,https://user:pw@example.org/a.pdf,2026-09-22,,\n"
        "2,article.pdf,article,https://example.org/b?session=abc,,,\n"
        "2,fake.pdf,article,,,,\n"
        "3,,article,,,unavailable,not held by library\n"
    )
    publications = {p: {"publication_id": f"pub:{p}", "doi": None, "title": f"Paper {p}"} for p in ("1", "2", "3")}
    result = run_intake(inbox, tmp_path / "papers", tmp_path, publications, set())
    states = {(x["pmid"], x["filename"]): x for x in result["records"]}
    assert (
        states[("1", "article.pdf")]["state"] == "INGESTED"
        and "pmid_in_text" in states[("1", "article.pdf")]["identity_evidence"]
    )
    assert states[("1", "article.pdf")]["source_url"] is None  # credentials never stored
    assert states[("2", "article.pdf")]["state"] == "IDENTITY_REVIEW_REQUIRED"  # a filename is not identity proof
    assert states[("2", "fake.pdf")]["state"] == "REJECTED"  # HTML masquerading as a PDF
    assert states[("3", "")]["state"] == "DISPOSITION_RECORDED"
    first = states[("1", "article.pdf")]["asset"]["sha256"]
    (inbox / "1" / "article.pdf").write_bytes(minimal_pdf("PMID 1 Paper about TEAD1 revised"))
    again = run_intake(inbox, tmp_path / "papers", tmp_path, publications, {first})
    revised = next(x for x in again["records"] if x["pmid"] == "1")
    assert (
        revised["asset"]["sha256"] != first and Path(tmp_path / states[("1", "article.pdf")]["asset"]["path"]).is_file()
    )
    assert safe_url("https://example.org/x?token=1") == (None, "source_url dropped: credential or session parameters")


def doc(pid: str, paragraphs: list[tuple[str, str]]) -> Document:
    """A synthetic parsed document: (section title, text) paragraphs under a Results/Introduction path."""
    from regkg.literature.parse import _Builder

    document = Document(pid, "europepmc_jats", hashlib.sha256(pid.encode()).hexdigest())
    build = _Builder(document)
    for i, (section, text) in enumerate(paragraphs):
        build.add("paragraph", text, f"/article/body/p[{i}]", section, section.lower())
    return document


LINEAGES = [
    {
        "cell_type": "AT1",
        "named_tfs": {"HGNC:1": ["TEAD1"]},
        "candidate_tfs": {"HGNC:1": ["TEAD1"]},
        "context_terms": [],
    },
    {
        "cell_type": "KRT",
        "named_tfs": {"HGNC:1": ["TEAD1"]},
        "candidate_tfs": {"HGNC:1": ["TEAD1"]},
        "context_terms": [],
    },
]


def screened(documents: dict[str, Document], relevance=None):
    meta = {d: {"publication_id": f"pub:{d}", "pmid": d, "role": "article", "route": "jats"} for d in documents}
    rel = relevance or {(d, c["cell_type"]): ("direct_context_candidate", "test") for d in documents for c in LINEAGES}
    return r.screen(documents, meta, LINEAGES, rel, r.gene_matcher({"HGNC:1": ["TEAD1"]})), meta


def test_every_paper_gets_candidates_even_beyond_the_old_global_top_ten():
    documents = {
        str(i): doc(str(i), [("Results", f"TEAD1 knockdown reduced CTGF in experiment {i}.")]) for i in range(15)
    }
    items, _ = screened(documents)
    assert {i["pmid"] for i in items if i["outcome"] == "bundle_candidate"} == set(documents)  # none cut at 10
    assert all(set(i["lineages"]) == {"AT1", "KRT"} for i in items)  # both lineages kept, no AT1 priority


def test_null_result_and_linked_context_without_the_tf_name_are_kept():
    document = doc(
        "1",
        [
            ("Results", "TEAD1 knockdown was performed in AT1 cells."),
            ("Results", "However, CTGF expression was unaffected by the knockdown."),
        ],
    )
    anchor = document.passages[0]
    bundle = r.build_bundle(anchor, document, "d1", {}, {"HGNC:1"})
    texts = [p["text"] for p in bundle["parts"]]
    assert "However, CTGF expression was unaffected by the knockdown." in texts  # no TF name, still linked
    assert "following_passage_continues_anchor" in bundle["context_reasons"] and "negation_cue" in bundle["flags"]


def test_broad_words_do_not_establish_manuscript_cell_types():
    from regkg.literature.screening import context_relevance, source_context

    bone = source_context("Alveolar bone mesenchymal stem cells in rabbit calvarial bone defects; TEAD1 expression")
    assert not bone["lung_tissue"] and context_relevance(bone, "AT1", False, [])[0] == "irrelevant_to_assigned_question"
    ligament = source_context("Anterior cruciate ligament-derived fibroblasts from rabbits")
    assert context_relevance(ligament, "Activated_Fibrotic_FBs", False, [])[0] == "irrelevant_to_assigned_question"
    cancer = source_context("Silencing of IRF8 expression in non-small cell lung cancer")
    assert context_relevance(cancer, "Alveolar_Macrophages", True, [])[0] == "potentially_transferable_mechanism"
    direct = source_context("TEAD1 knockdown in alveolar type I cells of the mouse lung")
    assert context_relevance(direct, "AT1", True, [])[0] == "direct_context_candidate"
    assert direct["species"] == {"mouse": ["mouse"]}  # reported species, independent of the manuscript association


def test_unique_bundles_are_counted_apart_from_query_associations():
    from regkg.workflows.corpus_prepare import _bundles

    documents = {"1": doc("1", [("Results", "TEAD1 knockdown reduced CTGF.")])}
    items, meta = screened(documents)
    queue = pd.DataFrame(
        {
            "query_id": ["q1", "q2", "q3"],
            "cell_type": ["AT1", "AT1", "KRT"],
            "tf_hgnc_id": ["HGNC:1"] * 3,
            "comparison_ids": [["c1"], ["c1"], ["c2"]],
        }
    )
    parse_of = {"1": {"state": "PARSED", "reasons": []}}
    payloads = _bundles(items, documents, meta, {}, {"HGNC:1"}, {}, queue, parse_of)
    assert len(payloads) == 1 and payloads[0]["query_associations"] == 3 and payloads[0]["lineages"] == ["AT1", "KRT"]


def test_serialized_limit_counts_cells_and_blocks_instead_of_truncating():
    from regkg.literature.parse import _Builder

    document = Document("p", "europepmc_jats", "a" * 64)
    build = _Builder(document)
    cells = [
        {"text": f"value {i}", "header": "log2 fold change (treated vs control)", "col_start": i} for i in range(80)
    ]
    build.add(
        "table_row",
        "TEAD1 | 1.5",
        "/t/tr[1]",
        "Results",
        "results",
        cells=cells,
        table_id="t",
        table_structure="header_grid_expanded",
        table_caption_status="present",
    )
    bundle = r.build_bundle(document.passages[0], document, "d", {}, {"HGNC:1"})
    assert (
        len(document.passages[0].text) < 100 and "serialized_size_exceeds_limit:anchor" in bundle["incomplete_reasons"]
    )
    assert bundle["parts"][0]["cells"] == cells  # evidence kept whole, not cut
    small = doc("q", [("Results", "TEAD1 binds. " * 10), ("Results", "This effect was lost. " + "x " * 3000)])
    fitted = r.build_bundle(small.passages[0], small, "d", {}, {"HGNC:1"})
    assert any(x.startswith("context_not_included:following_passage") for x in fitted["incomplete_reasons"])
    assert fitted["serialized_chars"] <= r.MAX_EVIDENCE_CHARS and fitted["part_count"] <= r.MAX_PARTS


def payload(pmid="1", version="europepmc_jats", refs=(), incomplete=()):
    return {
        "payload_id": f"payload:{pmid}",
        "publication_id": f"pub:{pmid}",
        "pmid": pmid,
        "text_version": version,
        "supplement_references": list(refs),
        "incomplete_reasons": list(incomplete),
        "anchor_passage_id": "a",
        "source_asset_sha256": "s",
        "lineages": ["AT1"],
    }


def test_readiness_is_fail_closed_and_unresolved_items_stay_unresolved():
    ok = {"state": "PARSED", "reasons": []}
    assert r.readiness(payload(), {"source_result_candidate"}, False, False, ok, []) == ("READY_FULL_TEXT", [])
    state, reasons = r.readiness(payload(), {"attribution_unresolved"}, False, False, ok, [])
    assert state == "NEEDS_SCREENING_REVIEW" and reasons[0].startswith("experimental_attribution_unresolved")
    assert (
        r.readiness(payload(version="pubmed_abstract"), {"abstract_summary"}, False, False, ok, [])[0] == "NEEDS_ASSET"
    )
    assert (
        r.readiness(payload(version="pubmed_abstract"), {"abstract_summary"}, True, False, ok, [])[0]
        == "READY_ABSTRACT_ONLY"
    )
    partial = {"state": "PARTIAL", "reasons": ["2 pages without text items"]}
    state, reasons = r.readiness(payload(), {"source_result_candidate"}, False, False, partial, [])
    assert state == "NEEDS_PARSE_REPAIR" and "2 pages without text items" in reasons[0]  # missing region exposed
    missing = r.readiness(
        payload(refs=["Supplementary Table S3"]), {"source_result_candidate"}, False, False, ok, ["missing_asset"]
    )
    assert missing[0] == "NEEDS_ASSET"
    unverified = r.readiness(
        payload(refs=["Fig. S2"]), {"source_result_candidate"}, False, False, ok, ["available_parsed_link_unverified"]
    )
    assert unverified[0] == "NEEDS_CONTEXT"  # a container file may not hold the cited item
    queued = r.investigations(
        [{**payload(refs=["Fig. S2"]), "readiness_state": unverified[0], "readiness_reasons": unverified[1]}],
        [{"payload_id": "payload:1", "local_asset_sha256": ["supp"]}],
    )
    assert queued[0]["status"] == "QUEUED_NOT_RUN" and queued[0]["permitted_asset_sha256"] == ["s", "supp"]


def test_missing_required_supplement_is_a_dependency_not_a_guess():
    from regkg.workflows.corpus_prepare import _dependencies

    p = {**payload(refs=["Supplementary Table 3"]), "payload_id": "b1"}
    supplements = [
        {
            "filename": "t2.xlsx",
            "label": "Supplementary Table 2",
            "state": "acquired",
            "url": "u2",
            "md5": None,
            "size": 1,
            "asset": {"sha256": "h2"},
        },
        {
            "filename": "t3.xlsx",
            "label": "Supplementary Table 3",
            "state": "not_in_oa_dataset",
            "url": None,
            "md5": None,
            "size": None,
        },
    ]
    deps, needed = _dependencies([p], {"1": {"supplements": supplements}}, {"supplements": [], "articles": {}})
    assert deps[0]["resolution"] == "resolved_by_label" and deps[0]["state"] == "missing_asset" and not needed
    none, _ = _dependencies([p], {"1": {"supplements": []}}, {"supplements": [], "articles": {}})
    assert none[0]["state"] == "missing_asset" and none[0]["files"] == []
    assert r.resolve_dependency("Fig. S4", supplements)[0] == "unresolved_reference"  # never matched by proximity


def test_ready_batches_respect_limits_rotate_lineages_and_queue_everything():
    payloads = [
        {
            "payload_id": f"p{i}",
            "publication_id": f"pub:{i}",
            "readiness_state": "READY_FULL_TEXT",
            "lineages": ["AT1"] if i % 3 else ["KRT"],
            "priority": 1.0 / (i + 1),
        }
        for i in range(30)
    ]
    batches = r.ready_batches(
        payloads, ["AT1", "KRT"], ReadyBatchLimits(max_papers=10, max_bundles=20, max_scheduled_calls=40)
    )
    assert sum(len(b["payload_ids"]) for b in batches) == 30
    assert all(len(b["papers"]) <= 10 and b["scheduled_calls"] <= 40 for b in batches)
    lineage = {p["payload_id"]: p["lineages"][0] for p in payloads}
    assert [lineage[x] for x in batches[0]["payload_ids"][:4]] == ["AT1", "KRT", "KRT", "AT1"]


def test_screening_and_priorities_replay_from_the_embedding_cache(tmp_path):
    class Backend:
        selector, cache_id, contract, identity = "fake:model", "fake", None, {"metadata": None}
        vector_identity = {"selector": "fake:model"}
        requests = 0

        class client:  # noqa: N801
            @staticmethod
            def embed_documents(texts):
                Backend.requests += 1
                return [[1.0, float(len(t) % 7)] for t in texts]

            @staticmethod
            def embed_query(text):
                return [1.0, 0.5]

    from regkg.config import EmbeddingSection

    documents = {str(i): doc(str(i), [("Results", f"TEAD1 knockdown changed target {i}.")]) for i in range(3)}
    items, _ = screened(documents)
    query = r.QueryTerms("q", 1, "HGNC:1", ["TEAD1"], None, [], ["knockdown"], [])
    runs = [
        r.prioritize(
            documents,
            [query],
            items,
            Backend,
            tmp_path,
            EmbeddingSection(max_chunk_chars=1000),
            60,
            1.5,
            0.75,
            log=lambda m: None,
        )
        for _ in range(2)
    ]
    assert runs[0][0] == runs[1][0] and runs[1][1]["cache_misses_before_encoding"] == 0
    assert screened(documents)[0] == items  # deterministic replay


def test_download_rows_use_only_returned_or_resolver_links(tmp_path):
    audits = acquire(tmp_path, Sources())
    row = {
        "publication_id": "pub:2",
        "pmid": "2",
        "doi": "10.1/p2",
        "title": "Paper 2",
        "year": 2020,
        "advisor_lists": '["a"]',
        "acquisition_priority": "high",
        "priority_reason": "r",
        "roles": "[]",
        "pubmed_url": "https://pubmed.ncbi.nlm.nih.gov/2/",
        "pmc_url": None,
        "doi_url": "https://doi.org/10.1/p2",
        "checked_sources": json.dumps(audits["2"]["checked_sources"]),
        "checked_at": "t",
        "acquisition_state": c.MANUAL_DOWNLOAD_NEEDED,
        "acquisition_reason": "not in the open-access subset",
        "download_need": "required_for_extraction",
        "returned_links": json.dumps([x for x in audits["2"]["links"] if x["kind"] == "returned_full_text"]),
    }
    downloads = c.manual_downloads(pd.DataFrame([row]), audits, {})
    links = json.loads(downloads.loc[0, "discovered_links"])
    assert [x["url"] for x in links] == ["https://europepmc.org/a2?pdf=render"]  # as returned, none constructed
    assert downloads.loc[0, "inbox_destination"] == "data/papers/manual_inbox/2/article.pdf"
    markdown = c.downloads_markdown("corpus-x", downloads, c.acquisition_failures({}), [])
    assert "Immediate downloads" in markdown and "proxy" in markdown


@pytest.mark.skipif(
    not files.docling_assets()["models"]["docling-project/docling-layout-heron"],
    reason="local Docling layout model not cached",
)
def test_docling_pdf_keeps_page_and_bbox_provenance():
    document, report, _ = files.parse_docling(
        "p:x", minimal_pdf("TEAD1 regulates CTGF in lung cells."), "a.pdf", "pdf", "manual_pdf", 10_000_000, 5, 120
    )
    assert report.state in {files.PARSED, files.PARTIAL} and report.parser_config["ocr"] is False
    passage = next(p for p in document.passages if "TEAD1" in p.text)
    assert passage.page == 1 and len(passage.bbox) == 4 and passage.locator.startswith("docling:#/texts/")
    assert document.extra_rules["pdf_rule"] == files.PDF_RULE


def test_tab_separated_txt_supplement_is_routed_as_a_table_not_prose(tmp_path):
    from regkg.config import PdfLimits, SpreadsheetLimits
    from regkg.literature.fetch import store_asset
    from regkg.workflows.corpus_parse import parse_supplement

    table = b"gene\tpeak\tscore\nTEAD1\tchr1:10-20\t5.2\nKLF5\tchr2:30-40\t3.1\nFOXA2\tchr3:1-9\t1.0\n"
    prose = b"Supplementary methods.\n\nCells were cultured as described.\n"
    results = {}
    for name, content in (("S1.txt", table), ("M1.txt", prose)):
        asset = store_asset(
            tmp_path / "papers", "pub:1", f"supplement-{name}", "supplement", content, "k", "u", tmp_path
        )
        [(document, report)] = parse_supplement(
            tmp_path / "work",
            tmp_path,
            tmp_path / "papers",
            "pub:1",
            asset.__dict__,
            PdfLimits(max_bytes=1, max_pages=1, timeout_seconds=1),
            SpreadsheetLimits(max_sheets=5, max_rows_per_sheet=100, max_columns=20),
        )
        results[name] = (document, report)
    table_doc, table_report = results["S1.txt"]
    assert (
        table_report["route"] == "delimited" and "text_file_detected_as_tab_delimited_table" in table_report["reasons"]
    )
    assert {p.kind for p in table_doc.passages} == {"supplement_table_header", "supplement_table_row"}
    assert results["M1.txt"][1]["route"] == "text"


def test_supplement_references_resolve_only_on_explicit_caption_or_filename_evidence():
    supplements = [
        {"filename": "mmc1.pdf", "label": None},
        {"filename": "JEM_TableS3.docx", "label": None},
        {"filename": "mmc2.xlsx", "label": None},
    ]
    captions = {"mmc1.pdf": "Document S1. Figures S1–S7 and Table S1"}
    assert r.resolve_dependency("Figure S5A", supplements, captions)[0] == "resolved_by_caption"
    assert r.resolve_dependency("Table S3", supplements, captions)[0] == "resolved_by_filename"
    assert r.resolve_dependency("Figure S9", supplements, captions)[0] == "unresolved_reference"  # beyond the range
    assert r._reference_key("Extended Data Fig. 2") != r._reference_key("Supplementary Fig. 2")


OBO = """format-version: 1.2
data-version: cl/releases/2000-01-01/cl-basic.owl

[Term]
id: CL:1
name: pulmonary alveolar type 1 cell
synonym: "membranous pneumocyte" EXACT []
synonym: "AT1" RELATED []
is_a: CL:9 ! epithelial cell of lung

[Term]
id: CL:2
name: wound-healing alveolar type 1 cell
is_a: CL:1

[Term]
id: CL:9
name: epithelial cell of lung

[Term]
id: CL:8
name: old term
is_obsolete: true
is_a: CL:1
"""


# --- P3 corpus corrections review: R1-R5 --------------------------------------------------------


def test_needed_supplements_and_archive_members_are_registered_and_partial_is_not_complete():
    from regkg.workflows.corpus_prepare import _asset_registry, _dependency_state, _merge_needed, _registry_check

    audits = {
        "1": {
            "pmid": "1",
            "publication_id": "pub:1",
            "assets": [],
            "supplements": [{"filename": "s1.zip", "state": "not_requested", "label": None}],
        }
    }
    asset = {
        "name": "supplement-s1.zip",
        "role": "supplement",
        "path": "p",
        "sha256": "zipsha",
        "bytes": 1,
        "source_cache_key": "k",
        "retrieved_at": "t",
        "url": "u",
    }
    needed = {
        "records": [
            {
                "results": [
                    {"pmid": "1", "publication_id": "pub:1", "filename": "s1.zip", "state": "acquired", "asset": asset}
                ]
            }
        ]
    }
    _merge_needed(audits, needed)
    assert audits["1"]["supplements"][0]["state"] == "acquired"
    parse_rows = [
        {"pmid": "1", "asset_sha256": "zipsha", "route": "zip", "state": "PARTIAL"},
        {
            "pmid": "1",
            "asset_sha256": "m1",
            "route": "xlsx",
            "state": "PARTIAL",
            "parent_asset": "zipsha",
            "archive_member": "t1.xlsx",
        },
        {
            "pmid": "1",
            "asset_sha256": "m2",
            "route": "text",
            "state": "PARSED",
            "parent_asset": "zipsha",
            "archive_member": "readme.txt",
        },
    ]
    needed["records"][0]["results"].append(
        {"pmid": "1", "publication_id": "pub:1", "filename": "big.zip", "state": "deferred_size_limit", "asset": None}
    )
    registry = _registry_check(needed, _asset_registry(audits, [], parse_rows), parse_rows, [])
    row, deferred = registry["files"]
    assert row["registered"] and row["parse_inventoried"] and row["members_registered"] == 2
    assert row["no_screening_items_reason"] == "parsed_no_candidate_tf_mention"
    assert deferred["archive_members"] == 0 and deferred["exclusion_reason"] == "deferred_size_limit"
    assert _dependency_state(audits["1"]["supplements"][0], parse_rows) == "available_partial"


def test_missing_label_reference_may_navigate_local_same_publication_sources_only():
    payload_a = {
        "payload_id": "b1",
        "publication_id": "pub:1",
        "pmid": "1",
        "anchor_passage_id": "a1",
        "source_asset_sha256": "article",
        "lineages": ["AT1"],
        "readiness_state": "NEEDS_CONTEXT",
        "readiness_reasons": ["cross_asset_context_required:Figure S2"],
    }
    payload_b = {**payload_a, "payload_id": "b2", "anchor_passage_id": "a2"}
    deps = [
        {"payload_id": "b1", "state": "unresolved_reference", "local_asset_sha256": []},
        {"payload_id": "b2", "state": "unresolved_reference", "local_asset_sha256": []},
    ]
    queue = r.investigations([payload_a, payload_b], deps, {"1": ["supp-local"], "2": ["other-paper"]})
    (entry,) = queue
    assert entry["permitted_asset_sha256"] == ["article", "supp-local"]  # never another paper's files
    assert entry["anchor_part_ids"] == ["a1", "a2"] and entry["bundle_ids"] == ["b1", "b2"]  # merged on dedupe
    assert "unverified" in entry["navigation_scope"]


class Onto:
    """A tiny synthetic ontology context (engineering fixture)."""

    record = {}
    strict = {
        "AT1": {"pulmonary alveolar type 1 cell": ["CL:1"]},
        "Alveolar_Macrophages": {"alveolar macrophage": ["CL:3"]},
        "KRT5neg_KRT17pos": {"aberrant basaloid": ["manuscript_label"]},
        "Activated_Fibrotic_FBs": {"activated fibroblast": ["manuscript_label"]},
    }
    broad = {"KRT5neg_KRT17pos": {"basaloid cell": ["grouping_phrase"]}}
    lung = {"lung": ["UBERON:1"], "lung tissue": ["UBERON:1"]}
    diseases = {"cancer": {"lung cancer": ["MONDO:2"]}}
    cl_all = {"alveolar type 1 fibroblast cell": ["CL:4"], "pulmonary alveolar type 1 cell": ["CL:1"]}


def test_review_probes_do_not_become_direct_context():
    from regkg.literature.screening import context_relevance, regulatory_content, source_context

    probes = [
        ("AT1", "TEAD1 regulates alveolar type 1 fibroblast cells in lung tissue."),
        ("KRT5neg_KRT17pos", "KLF5 regulates basaloid cells in lung cancer."),
        ("Activated_Fibrotic_FBs", "We investigated activated fibroblasts in skin injury and lung epithelial cells."),
    ]
    for cell_type, text in probes:
        profile = source_context(text, Onto())
        assert context_relevance(profile, cell_type, regulatory_content(text), [])[0] != "direct_context_candidate"
    at1 = source_context("TEAD1 regulates alveolar type 1 fibroblast cells in lung tissue.", Onto())
    assert at1["cell_types"]["AT1"]["overridden_by_longer_cl_name"]  # longer CL name wins
    positive = source_context("KLF5 is expressed by aberrant basaloid cells in the IPF lung.", Onto())
    assert context_relevance(positive, "KRT5neg_KRT17pos", True, [])[0] == "direct_context_candidate"


def test_full_text_result_survives_a_misleading_abstract_and_case_variants_stay_unresolved():
    documents = {
        "1": doc(
            "1",
            [
                ("Results", "Egr2 deletion in alveolar macrophages of the lung reduced CD11c."),
                ("Results", "TEAD1 knockdown in lung AT1 cells had no effect on AGER."),
            ],
        )
    }
    meta = {"1": {"publication_id": "pub:1", "pmid": "1", "role": "article", "route": "jats"}}
    lineages = [
        {"cell_type": "Alveolar_Macrophages", "candidate_tfs": {"HGNC:2": ["EGR2"]}},
        {"cell_type": "AT1", "candidate_tfs": {"HGNC:1": ["TEAD1"]}},
    ]
    irrelevant = {
        ("1", c): ("irrelevant_to_assigned_question", "abstract says bone") for c in ("Alveolar_Macrophages", "AT1")
    }
    items = r.screen(
        documents, meta, lineages, irrelevant, r.gene_matcher({"HGNC:2": ["EGR2"], "HGNC:1": ["TEAD1"]}), Onto()
    )
    outcomes = {i["tf_hgnc_id"]: i for i in items}
    assert outcomes["HGNC:2"]["outcome"] == "bundle_candidate"  # the abstract heuristic alone never excludes
    assert outcomes["HGNC:2"]["match_kinds"] == ["case_variant_species_unresolved"]
    assert outcomes["HGNC:1"]["outcome"] == "bundle_candidate"  # a null result is kept
    bundle = {
        **r.build_bundle(documents["1"].passages[0], documents["1"], "1", {}, {"HGNC:2"}),
        "anchor_passage_id": documents["1"].passages[0].passage_id,
    }
    identity = r.tf_identity(bundle, {"HGNC:2": {"case_variant_species_unresolved"}})
    assert identity["status"] == "case_variant_species_unresolved"
    assert identity["scorable_as_resolved_human"] is False and identity["unresolved_tfs"] == ["HGNC:2"]
    assert identity["per_tf"]["HGNC:2"]["source_identifiers"] == []  # no human ID invented


def test_ambiguous_source_identifier_blocks_but_a_species_conflict_is_only_flagged():
    ok = {"state": "PARSED", "reasons": []}
    ambiguous = {**payload(), "tf_identity": {"status": "ambiguous_source_identifier"}}
    assert r.readiness(ambiguous, {"source_result_candidate"}, False, False, ok, [])[0] == "NEEDS_SCREENING_REVIEW"
    conflict = {
        **payload(),
        "tf_identity": {"status": "name_conflicts_with_source_identifier", "scorable_as_resolved_human": False},
    }
    assert r.readiness(conflict, {"source_result_candidate"}, False, False, ok, [])[0] == "READY_FULL_TEXT"


def test_download_report_separates_unidentified_references_navigation_and_identified_files():
    from regkg.workflows.corpus_prepare import _downloads

    audits = {
        "1": {
            "pmid": "1",
            "publication_id": "pub:1",
            "supplements": [],
            "links": [],
            "errors": [],
            "is_open_access": True,
        },
        "2": {
            "pmid": "2",
            "publication_id": "pub:2",
            "links": [],
            "errors": [],
            "is_open_access": True,
            "supplements": [{"filename": "s.xlsx", "asset": {"sha256": "x"}}],
        },
    }
    deps = [
        {
            "dependency_id": "d1",
            "pmid": "1",
            "publication_id": "pub:1",
            "payload_id": "b1",
            "reference": "Table S1",
            "state": "unresolved_reference",
            "resolution": "unresolved_reference",
            "files": [],
            "local_asset_sha256": [],
        },
        {
            "dependency_id": "d2",
            "pmid": "1",
            "publication_id": "pub:1",
            "payload_id": "b2",
            "reference": "Figure S4",
            "state": "unresolved_reference",
            "resolution": "unresolved_reference",
            "files": [],
            "local_asset_sha256": [],
        },
        {
            "dependency_id": "d3",
            "pmid": "1",
            "publication_id": "pub:1",
            "payload_id": "b3",
            "reference": "Table S9",
            "state": "missing_asset",
            "resolution": "resolved_by_label",
            "local_asset_sha256": [],
            "files": [{"filename": "t9.xlsx", "url": None}],
        },
        {
            "dependency_id": "d4",
            "pmid": "1",
            "publication_id": "pub:1",
            "payload_id": "b4",
            "reference": "Table 9",
            "state": "missing_asset",
            "resolution": "resolved_by_caption",
            "local_asset_sha256": [],
            "files": [{"filename": "t9.xlsx", "url": None}],
        },
        {
            "dependency_id": "d5",
            "pmid": "2",
            "publication_id": "pub:2",
            "payload_id": "b5",
            "reference": "Fig. S1",
            "state": "unresolved_reference",
            "resolution": "unresolved_reference",
            "files": [],
            "local_asset_sha256": [],
        },
    ]
    downloads, tasks = _downloads(None, audits, {}, deps, None, {})
    unknown = downloads[downloads["group"] == "unidentified_reference"]
    assert len(unknown) == 2 and unknown["request_id"].is_unique and unknown["inbox_destination"].isna().all()
    missing = downloads[downloads["group"] == "identified_missing_file"]
    assert len(missing) == 1 and sorted(json.loads(missing.iloc[0]["affected_bundles"])) == ["b3", "b4"]
    assert list(tasks["pmid"]) == ["2"]  # a local candidate exists: navigation, not a download


def test_budget_charges_downstream_once_and_shows_the_assembly_ceiling():
    assumptions = BudgetAssumptions(
        chars_per_token=4,
        extractor_overhead_tokens=100,
        verifier_overhead_tokens=100,
        extractor_output_tokens_typical=10,
        extractor_output_tokens_cap=20,
        verifier_output_tokens_typical=10,
        verifier_output_tokens_cap=20,
        reasoning_tokens_typical=5,
        reasoning_tokens_cap=10,
        max_attempts_per_call=3,
        schema_preflight_calls=2,
        planning_reserve=0.25,
        illustrative_rates={"x": (1.0, 2.0)},
        assembly={
            "input_tokens_cap_per_call": 8000,
            "preflight_calls_per_batch": 1,
            "low": {
                "calls": 1,
                "first_call_input_tokens": 3000,
                "history_growth_per_call_tokens": 0,
                "output_tokens_per_call": 400,
            },
            "typical": {
                "calls": 2,
                "first_call_input_tokens": 3500,
                "history_growth_per_call_tokens": 2500,
                "output_tokens_per_call": 600,
            },
            "conservative": {
                "calls": 4,
                "first_call_input_tokens": 4000,
                "history_growth_per_call_tokens": 2000,
                "output_tokens_per_call": 1000,
            },
        },
    )
    blocked = {
        "payload_id": "b",
        "serialized_chars": 20000,
        "readiness_state": "NEEDS_CONTEXT",
        "lineages": ["AT1"],
        "relevance": {},
    }
    queue = [{"investigation_id": "i1", "bundle_ids": ["b"]}, {"investigation_id": "i2", "bundle_ids": ["b"]}]
    estimate = r.budget([blocked], [], queue, assumptions, ["AT1"])
    ceiling = estimate["scenarios"]["ceiling"]
    assert estimate["counts"]["repaired_bundles_bound"] == 1  # shared repair charged once
    assert ceiling["assembly_plus_downstream"]["calls"] - ceiling["assembly_only_no_result"]["calls"] == 2 * 3
    only = r.budget([blocked], [], queue[:1], assumptions, ["AT1"])["scenarios"]["ceiling"]["assembly_only_no_result"]
    assert only["input"] >= 4 * 8000 and only["output"] >= 4 * 1000  # the contract's configured maximum
    assert estimate["counts"]["hypothetical_split_children"] == 4  # 20,000 chars -> <= 6,000-char children


def pinned_scope(tmp_path, alias: str, corrupt: bool = False):
    import hashlib

    from regkg.resources import Fetcher, store_snapshot

    cl = OBO.replace('"membranous pneumocyte" EXACT', f'"{alias}" EXACT')
    uberon = "format-version: 1.2\n\n[Term]\nid: UBERON:0002048\nname: lung\n"
    mondo = "format-version: 1.2\n\n[Term]\nid: MONDO:1\nname: idiopathic pulmonary fibrosis\n"
    releases = {}
    for name, text, file in (("cl", cl, "cl.obo"), ("uberon", uberon, "u.obo"), ("mondo", mondo, "m.obo")):
        store_snapshot(
            tmp_path / "ext" / "ontologies" / name / "2000-01-01",
            file,
            text.encode(),
            "u",
            Fetcher(),
            expected_md5_hex=None,
            metadata={},
            response_headers={},
        )
        releases[name] = {
            "version": "2000-01-01",
            "release": "2000-01-01",
            "file": file,
            "sha256": hashlib.sha256(text.encode()).hexdigest(),
        }
    if corrupt:
        (tmp_path / "ext" / "ontologies" / "cl" / "2000-01-01" / "cl.obo").write_text(cl + "\n")
    scope = {
        "schema_version": 1,
        "mapping_version": "t",
        "status": "proposed",
        "releases": releases,
        "cell_types": {"AT1": {"exact_cl": ["CL:1"], "grouping_parents": ["CL:9"]}},
        "anatomy": {"lung": ["UBERON:0002048"]},
        "diseases": {"pulmonary_fibrosis": ["MONDO:1"]},
    }
    path = tmp_path / "scope.yaml"
    path.write_text(yaml.safe_dump(scope))
    return path, tmp_path / "ext"


def test_ontology_identity_changes_with_same_count_synonym_edits_and_rejects_bad_resources(tmp_path):
    from regkg.literature.ontology import build_context
    from regkg.resources import ResourceError

    first = build_context(*pinned_scope(tmp_path / "a", "membranous pneumocyte")).record
    again = build_context(*pinned_scope(tmp_path / "b", "membranous pneumocyte")).record
    edited = build_context(*pinned_scope(tmp_path / "c", "squamous pneumocyte")).record
    assert first["phrase_tables_sha256"] == again["phrase_tables_sha256"]  # unchanged replay reuses
    assert first["phrases"] == edited["phrases"] and first["phrase_tables_sha256"] != edited["phrase_tables_sha256"]
    with pytest.raises(ResourceError, match="no longer matches"):
        build_context(*pinned_scope(tmp_path / "d", "membranous pneumocyte", corrupt=True))
    scope, _ = pinned_scope(tmp_path / "e", "membranous pneumocyte")
    with pytest.raises(FileNotFoundError, match="ontology-setup"):
        build_context(scope, tmp_path / "missing")


def test_a_phrase_shared_by_another_cl_class_is_demoted_to_grouping_level(tmp_path):
    from regkg.literature.ontology import build_context

    obo = OBO + '\n[Term]\nid: CL:7\nname: other cell\nsynonym: "membranous pneumocyte" EXACT []\n'
    scope, ext = pinned_scope(tmp_path, "membranous pneumocyte")
    (ext / "ontologies" / "cl" / "2000-01-01" / "cl.obo").unlink()
    import hashlib

    from regkg.resources import Fetcher, store_snapshot

    store_snapshot(
        ext / "ontologies" / "cl" / "2000-01-01",
        "cl.obo",
        obo.encode(),
        "u",
        Fetcher(),
        expected_md5_hex=None,
        metadata={},
        response_headers={},
    )
    data = yaml.safe_load(scope.read_text())
    data["releases"]["cl"]["sha256"] = hashlib.sha256(obo.encode()).hexdigest()
    scope.write_text(yaml.safe_dump(data))
    context = build_context(scope, ext)
    assert "membranous pneumocyte" not in context.strict["AT1"]
    assert context.broad["AT1"]["membranous pneumocyte"] == ["CL:1", "CL:7"]  # all IDs kept, none chosen


def test_jats_break_separates_words():
    from regkg.literature.xml import element_text, parse_xml

    assert element_text(parse_xml(b"<th>Gene<break/>TF Rank</th>")) == "Gene TF Rank"


def test_synonym_followup_never_reacquires_p2_papers_and_needs_tf_plus_context():
    from types import SimpleNamespace

    from regkg.literature.search import SourceRecord
    from regkg.workflows import corpus as c

    queue = c.synonym_queue(
        {
            "queries": [
                {
                    "cell_type": "AT1",
                    "tf_hgnc_id": "HGNC:1",
                    "tf_terms": ["TEAD1"],
                    "context_phrases": ["type i pneumocyte"],
                    "phrase_ontology_ids": {"type i pneumocyte": ["CL:1"]},
                }
            ]
        }
    )
    records = [
        SourceRecord("pubmed", 1, "k", pmid="10", title="TEAD1 in type I pneumocyte spreading"),
        SourceRecord("pubmed", 2, "k", pmid="11", title="TEAD1 in type I pneumocytes, found by P2 too"),
        SourceRecord("pubmed", 3, "k", pmid="12", title="Hippo signalling in type I pneumocyte"),
        SourceRecord("pubmed", 4, "k", pmid="13", title="TEAD1 in cardiomyocytes"),
    ]
    result = {
        "records": [
            {
                "results": [
                    {"query_id": queue[0]["query_id"], "status": "success", "records": [x.__dict__ for x in records]}
                ]
            }
        ]
    }
    ctx = SimpleNamespace(
        advisors=pd.DataFrame({"pmid": ["99"]}),
        scope=SimpleNamespace(lineages=[SimpleNamespace(cell_type="AT1", context_terms=["lung"])]),
    )
    p2 = pd.DataFrame({"publication_id": ["pub:other"], "pmid": ["11"]})
    frame, order = c.synonym_publications(ctx, queue, result, p2)
    tiers = dict(zip(frame["pmid"], frame["tier"], strict=True))
    assert tiers == {
        "10": "tier1_acquired",
        "11": "already_in_p2_discovery",
        "12": "metadata_only_no_exact_tf_in_title_abstract",
        "13": "tier2_deferred_no_lineage_context_in_title_abstract",
    }
    assert [pmid for _, pmid in order] == ["10"]
    assert json.loads(queue[0]["terms_json"]) == {"qualifiers": ["type i pneumocyte"], "target": [], "tf": ["TEAD1"]}


def _mention(tf, status, start, text, resolution="candidate_name_match_species_unknown", hgnc=None, taxon=None):
    return {
        "candidate_hgnc_id": None if hgnc else tf,
        "hgnc_id": hgnc,
        "identity_status": status,
        "resolution": resolution,
        "species_taxon": taxon,
        "start": start,
        "end": start + len(text),
        "text": text,
    }


def _mixed_bundle(second_status: str) -> dict:
    """A bundle whose anchor names TF A (source-annotated human) and TF B, anchored by two screening
    groups that merge into one payload (synthetic)."""
    from regkg.workflows.corpus_prepare import _bundles

    text = "YAP1 and TEAD1 bind the AGER promoter in AT1 cells of the lung."
    document = doc("1", [("Results", text)])
    passage = document.passages[0]
    mentions = {
        passage.passage_id: [
            _mention("HGNC:A", "consistent_with_source_identifier", 0, "YAP1"),
            _mention("HGNC:A", "source_identifier", 0, "YAP1", "resolved_ncbi_gene_human", "HGNC:A", "9606"),
            _mention("HGNC:B", second_status, 9, "TEAD1"),
        ]
    }
    items = [
        {
            "item_id": f"i{tf}",
            "passage_id": passage.passage_id,
            "tf_hgnc_id": tf,
            "outcome": "bundle_candidate",
            "attribution": "source_result_candidate",
            "match_kinds": ["exact_symbol"],
            "lineages": {"AT1": {"relevance": "direct_context_candidate", "reason": "x"}},
        }
        for tf in ("HGNC:A", "HGNC:B")
    ]
    queue = pd.DataFrame(
        {
            "query_id": ["qA", "qB"],
            "cell_type": ["AT1", "AT1"],
            "tf_hgnc_id": ["HGNC:A", "HGNC:B"],
            "comparison_ids": [["c"], ["c"]],
        }
    )
    meta = {"doc1": {"publication_id": "pub:1", "pmid": "1", "role": "article"}}
    parse_of = {"doc1": {"state": "PARSED", "reasons": []}}
    (payload,) = _bundles(items, {"doc1": document}, meta, mentions, {"HGNC:A", "HGNC:B"}, {}, queue, parse_of)
    return payload


def test_a_resolved_tf_never_upgrades_a_name_only_tf_in_the_same_bundle():
    identity = _mixed_bundle("candidate_hint_only")["tf_identity"]
    assert identity["scorable_as_resolved_human"] is False
    assert identity["resolved_tfs"] == ["HGNC:A"] and identity["unresolved_tfs"] == ["HGNC:B"]
    assert identity["per_tf"]["HGNC:B"]["status"] == "name_only_species_unknown"
    assert identity["per_tf"]["HGNC:A"]["source_identifiers"] == [
        {"hgnc_id": "HGNC:A", "resolution": "resolved_ncbi_gene_human", "species_taxon": "9606"}
    ]


def test_a_resolved_tf_does_not_mask_a_conflicting_tf_and_manifests_carry_per_tf_identity():
    payload = _mixed_bundle("conflicting_source_identifier")
    identity = payload["tf_identity"]
    assert identity["status"] == "name_conflicts_with_source_identifier"
    assert identity["scorable_as_resolved_human"] is False
    assert identity["per_tf"]["HGNC:A"]["scorable_as_resolved_human"] is True
    assert payload["identity_flags"] == ["name_conflicts_with_source_identifier"]
    payload.update(readiness_state="READY_FULL_TEXT", readiness_reasons=[], serialized_chars=10)
    manifest = r.batch_manifest(
        {"batch_index": 1, "payload_ids": [payload["payload_id"]]}, {payload["payload_id"]: payload}, {}
    )
    assert manifest["items"][0]["tf_identity"]["per_tf"]["HGNC:B"]["status"] == "name_conflicts_with_source_identifier"


def test_pilot_respects_caps_excludes_unresolved_identity_and_reports_missing_lineages():
    from regkg.workflows import corpus_pilot as pilot

    def bundle(i, pmid, lineage, tf, passage="unresolved", paper="potentially_transferable_mechanism", ok=True):
        return {
            "payload_id": f"b{i:03d}",
            "pmid": pmid,
            "readiness_state": "READY_FULL_TEXT",
            "priority": 1.0 / (i + 1),
            "attributions": ["source_result_candidate"],
            "tf_identity": {
                "resolved_tfs": [tf] if ok else [],
                "per_tf": {tf: {"scorable_as_resolved_human": ok}},
            },
            "associations": [
                {"cell_type": lineage, "tf_hgnc_id": tf, "relevance": paper, "passage_relevance": passage}
            ],
        }

    payloads = [bundle(i, str(i % 14), "AT1" if i % 2 else "AM", f"T{i % 7}") for i in range(60)]
    payloads.append(bundle(98, "99", "KRT", "K1", ok=False))  # a name-only TF never enters the pool
    payloads.append(bundle(99, "98", "FB", "F1", paper="biological_background_resource"))
    named = {"AT1": {"T1"}, "AM": {"T2"}, "KRT": {"K1"}, "FB": {"F1"}}
    mentions = {p["payload_id"]: dict.fromkeys(named, "broad") for p in payloads}
    mentions["b059"]["AT1"] = None  # an anchor that never names the cell class cannot be selected
    options, excluded = pilot.candidates(payloads, named, mentions)
    assert "b059" not in {o["payload_id"] for o in options["AT1"]}
    assert excluded["AT1"]["anchor_does_not_name_lineage_cell_class"] == 1
    assert excluded["KRT"] == {"no_tf_consistent_with_source_identifier": 1}
    chosen = pilot.select(options)
    assert len(chosen) <= 20 and len({x["pmid"] for x in chosen}) <= 10
    assert max(sum(x["pmid"] == p for x in chosen) for p in {x["pmid"] for x in chosen}) <= 3
    assert {x["lineage"] for x in chosen} == {"AT1", "AM"}
    report = pilot.coverage(options, chosen, named)
    assert report["KRT"]["gaps"] == ["no_eligible_bundle"] and report["FB"]["gaps"] == ["no_eligible_bundle"]
    assert report["KRT"]["named_regulators_without_eligible_bundle"] == ["K1"]


def test_pilot_judges_identity_per_association_not_per_bundle():
    from regkg.workflows import corpus_pilot as pilot

    payload = {
        "payload_id": "b1",
        "pmid": "1",
        "readiness_state": "READY_FULL_TEXT",
        "priority": 1.0,
        "attributions": ["source_result_candidate"],
        "tf_identity": {
            "resolved_tfs": ["A"],
            "per_tf": {"A": {"scorable_as_resolved_human": True}, "B": {"scorable_as_resolved_human": False}},
        },
        "associations": [
            {
                "cell_type": "AM",
                "tf_hgnc_id": tf,
                "relevance": "direct_context_candidate",
                "passage_relevance": "direct_context_candidate",
            }
            for tf in ("A", "B")
        ],
    }
    options, excluded = pilot.candidates([payload], {"AM": set()}, {"b1": {"AM": "strict"}})
    assert options["AM"][0]["tf_hgnc_ids"] == ["A"]  # B is never carried in on A's identity
    assert excluded["AM"] == {"association_tf_not_consistent_with_source_identifier": 1}


# --- P3 scope closure ---------------------------------------------------------------------------


def _closure_artifact(tmp_path):
    """A tiny synthetic corpus artifact with one item of every frontier kind."""
    from regkg.workflows import scope_closure as sc

    artifact = tmp_path / "litcorpus-test"
    (artifact / "x").parent.mkdir(parents=True, exist_ok=True)

    def payload(pid, state, relevance, attribution, tfs):
        return {
            "payload_id": pid,
            "pmid": "1",
            "readiness_state": state,
            "readiness_reasons": [],
            "attributions": [attribution],
            "tf_hgnc_ids": list(tfs),
            "priority": 1.0,
            "serialized_chars": 10,
            "publication_id": "pub:1",
            "document_id": "d",
            "source_asset_sha256": "a",
            "lineages": ["AT1"],
            "relevance": {},
            "anchor_passage_id": "p1",
            "parts": [
                {
                    "passage_id": "p1",
                    "locator": "loc",
                    "asset_sha256": "a",
                    "text": "TEAD1 knockout in mouse lung fibroblasts changed collagen.",
                    "mentions": [],
                }
            ],
            "tf_identity": {
                "per_tf": {
                    t: {
                        "status": "consistent_with_source_identifier" if t == "HGNC:1" else "name_only_species_unknown",
                        "scorable_as_resolved_human": t == "HGNC:1",
                    }
                    for t in tfs
                },
                "resolved_tfs": [t for t in tfs if t == "HGNC:1"],
            },
            "associations": [
                {"cell_type": "AT1", "tf_hgnc_id": t, "relevance": relevance, "passage_relevance": relevance}
                for t in tfs
            ],
        }

    payloads = [
        payload(
            "b:ready", "READY_FULL_TEXT", "direct_context_candidate", "source_result_candidate", ["HGNC:1", "HGNC:2"]
        ),
        payload("b:background", "READY_FULL_TEXT", "biological_background_resource", "abstract_summary", ["HGNC:1"]),
        payload("b:context", "NEEDS_CONTEXT", "unresolved", "source_result_candidate", ["HGNC:1"]),
        payload("b:asset", "NEEDS_ASSET", "unresolved", "source_result_candidate", ["HGNC:9"]),
    ]
    (artifact / "bundles_all.jsonl").write_text("".join(json.dumps(p) + "\n" for p in payloads))
    (artifact / "source_bundles_ready.jsonl").write_text("")
    (artifact / "manifest.json").write_text("{}")
    (artifact / "ontology_scope.json").write_text(json.dumps({"mapping_sha256": "abc"}))
    (artifact / "context_investigations.jsonl").write_text(
        json.dumps({"investigation_id": "inv:1", "bundle_ids": ["b:context"]}) + "\n"
    )
    pd.DataFrame(
        [
            {
                "kind": "bundle",
                "id": p["payload_id"],
                "pmid": "1",
                "state": p["readiness_state"],
                "ready_batch": None,
                "reasons": json.dumps(["cross_asset_context_required:S1"]),
                "lineages": json.dumps(["AT1"]),
            }
            for p in payloads
        ]
        + [
            {
                "kind": "context_investigation",
                "id": "inv:1",
                "pmid": "1",
                "state": "queued",
                "ready_batch": None,
                "reasons": None,
                "lineages": json.dumps(["AT1"]),
            },
            {
                "kind": "table_region_interpretation",
                "id": "screen:1",
                "pmid": "1",
                "state": "review",
                "ready_batch": None,
                "reasons": None,
                "lineages": json.dumps(["AT1"]),
            },
            {
                "kind": "deferred_discovery_paper",
                "id": "pub:2",
                "pmid": "2",
                "state": "tier2",
                "ready_batch": None,
                "reasons": None,
                "lineages": json.dumps(["AT1"]),
            },
            {
                "kind": "download:unidentified_reference",
                "id": "download:nav",
                "pmid": "1",
                "state": "NEEDS_IDENTIFICATION",
                "ready_batch": None,
                "reasons": None,
                "lineages": json.dumps(["AT1"]),
            },
            {
                "kind": "download:unidentified_reference",
                "id": "download:unknown",
                "pmid": "1",
                "state": "NEEDS_IDENTIFICATION",
                "ready_batch": None,
                "reasons": None,
                "lineages": json.dumps(["AT1"]),
            },
            {
                "kind": "download:article_lower_priority_background",
                "id": "download:bg",
                "pmid": "3",
                "state": "NEEDED",
                "ready_batch": None,
                "reasons": None,
                "lineages": json.dumps([]),
            },
        ]
    ).to_parquet(artifact / "remaining_work.parquet", index=False)
    pd.DataFrame(
        [
            {
                "request_id": "download:nav",
                "group": "unidentified_reference",
                "pmid": "1",
                "title": "t",
                "item": "Table S1",
                "why": "w",
                "verified_link": None,
                "doi_url": None,
                "pubmed_url": "u",
                "inbox_destination": None,
                "affected_bundles": json.dumps(["b:context"]),
                "state": "NEEDS_IDENTIFICATION",
            },
            {
                "request_id": "download:unknown",
                "group": "unidentified_reference",
                "pmid": "1",
                "title": "t",
                "item": "Figure S9",
                "why": "w",
                "verified_link": None,
                "doi_url": None,
                "pubmed_url": "u",
                "inbox_destination": None,
                "affected_bundles": json.dumps([]),
                "state": "NEEDS_IDENTIFICATION",
            },
            {
                "request_id": "download:bg",
                "group": "article_lower_priority_background",
                "pmid": "3",
                "title": "t",
                "item": "article PDF",
                "why": "w",
                "verified_link": "l",
                "doi_url": None,
                "pubmed_url": "u",
                "inbox_destination": "data/papers/manual_inbox/3/article.pdf",
                "affected_bundles": json.dumps([]),
                "state": "NEEDED",
            },
        ]
    ).to_csv(artifact / "manual_downloads.csv", index=False)
    pd.DataFrame(
        [
            {
                "task_id": "download:nav",
                "pmid": "1",
                "publication_id": "pub:1",
                "reference": "Table S1",
                "resolution": "unresolved",
                "payload_id": "b:context",
                "candidate_local_assets": "[]",
                "kind": "navigate",
            }
        ]
    ).to_csv(artifact / "source_navigation_tasks.csv", index=False)
    pd.DataFrame([{"item_id": "screen:1", "tf_hgnc_id": "HGNC:1", "pmid": "1"}]).to_parquet(
        artifact / "screening_ledger.parquet", index=False
    )
    pd.DataFrame(
        [
            {
                "publication_id": "pub:2",
                "pmid": "2",
                "lineage_associations": json.dumps({"AT1": {"tfs_named": ["HGNC:1"]}}),
            }
        ]
    ).to_parquet(artifact / "discovery_coverage.parquet", index=False)
    return sc, artifact, {p["payload_id"]: p for p in payloads}


def test_every_frontier_item_gets_exactly_one_disposition_and_totals_reconcile(tmp_path):
    sc, artifact, _ = _closure_artifact(tmp_path)
    frame = sc.dispositions(artifact, {"AT1": {"HGNC:1"}})
    assert len(frame) == 10 and frame["id"].is_unique and frame["disposition"].notna().all()
    by_id = frame.set_index("id")["disposition"].to_dict()
    assert by_id["b:ready"] == sc.EXTRACTION
    # a provisional background/irrelevant screening label is not a reviewed scientific exclusion
    assert by_id["b:background"] == sc.UNRESOLVED
    assert frame.set_index("id").loc["b:background", "scientific_relevance"] == "provisional_not_reviewed"
    assert by_id["b:context"] == sc.INVESTIGATION and by_id["inv:1"] == sc.INVESTIGATION
    assert by_id["b:asset"] == sc.REQUIRED and by_id["screen:1"] == sc.REQUIRED
    assert by_id["download:nav"] == sc.INVESTIGATION  # a local candidate exists: navigate before requesting
    assert by_id["download:unknown"] == sc.REQUIRED
    assert by_id["pub:2"] == sc.UNRESOLVED  # a deferred paper is never silently dropped
    # the unrequested background article closes only the acquisition action
    background = frame.set_index("id").loc["download:bg"]
    assert background["disposition"] == sc.BACKGROUND and background["closes"] == "acquisition_action_only"
    assert background["scientific_relevance"] == "provisional_not_reviewed"
    assert int(frame["open"].sum()) == 9
    # linked work is recorded, never summed: the investigation and its bundle point at each other
    assert json.loads(frame.set_index("id").loc["b:context", "linked_to"]) == ["inv:1"]
    assert json.loads(frame.set_index("id").loc["inv:1", "linked_to"]) == ["b:context"]
    assert set(frame["unit_group"]) == {
        "evidence_bundles",
        "investigations_linked_to_blocked_bundles",
        "table_regions_pending_interpretation",
        "papers_pending_relevance_review",
        "acquisition_or_identification_requests",
    }
    critical = frame.set_index("id")["names_manuscript_regulator"].to_dict()
    assert critical["b:ready"] and not critical["b:asset"]  # HGNC:9 is not an AT1 named regulator


def test_closure_rejects_spans_and_associations_a_bundle_does_not_contain(tmp_path):
    sc, _, payloads = _closure_artifact(tmp_path)
    good = {
        "rule": "r",
        "corpus_key": "k",
        "audit": [
            {
                "payload_id": "b:ready",
                "species": ["mouse", "mouse lung fibroblasts"],
                "cell": "unknown",
                "attribution": "unknown",
                "assay": "unknown",
            }
        ],
        "selection": [
            {
                "payload_id": "b:ready",
                "associations": [["AT1", "HGNC:1"]],
                "purpose": "genuine_evidence",
                "coverage": "exact_population",
            }
        ],
    }
    assert sc.check_spans(good, payloads) == []
    bad = json.loads(json.dumps(good))
    bad["audit"][0]["species"] = ["human", "human donor lungs"]  # not in the bundle parts
    bad["selection"][0]["associations"] = [["Alveolar_Macrophages", "HGNC:1"]]
    problems = {p["problem"] for p in sc.check_spans(bad, payloads)}
    assert "species_span_not_in_bundle_parts" in problems
    assert "association_absent:Alveolar_Macrophages/HGNC:1" in problems


def test_selection_manifest_keeps_unselected_co_mentioned_identities(tmp_path):
    sc, _, payloads = _closure_artifact(tmp_path)
    document = {
        "rule": "r",
        "corpus_key": "k",
        "audit": [
            {
                "payload_id": "b:ready",
                "species": "unknown",
                "cell": "unknown",
                "attribution": "unknown",
                "assay": "unknown",
            }
        ],
        "selection": [
            {
                "payload_id": "b:ready",
                "associations": [["AT1", "HGNC:1"]],
                "purpose": "genuine_evidence",
                "coverage": "exact_population",
                "must_reject": ["human evidence"],
            }
        ],
    }
    manifest = sc.selection_manifest(document, payloads, {})
    validation = manifest["items"][0]["validation"]
    assert validation["selected_associations"][0]["tf_identity"]["scorable_as_resolved_human"] is True
    assert validation["co_mentioned_tf_identity"]["HGNC:2"]["status"] == "name_only_species_unknown"
    assert validation["expected"]["must_reject"] == ["human evidence"]
    assert manifest["status"] == "proposed_validation_selection_pending_lead_review"


# --- P3 source-gap completion -------------------------------------------------------------------


def _big_table(rows: int = 5000) -> bytes:
    body = ["cellType\tgene\tlog1p_FC\tpct.inGroup\n"]
    body += [f"AT2\tGENE{i}\t0.1\t0.5\n" for i in range(rows)]
    body.append("Aberrant_Basaloid\tTEAD1\t0.42\t0.61\n")  # the row the old 2,000-row limit cut off
    return "".join(body).encode()


def test_delimited_scan_covers_every_row_and_materializes_matches_past_the_old_limit():
    from regkg.literature.files import RowSelector, parse_delimited

    selector = RowSelector("test", frozenset({"TEAD1"}), frozenset({"aberrant_basaloid"}))
    document, report = parse_delimited("pub:1", _big_table(), "S3.txt", "\t", 2000, 60, selector)
    result = report.to_json()
    assert result["components"]["rows_total"] == 5002 and result["components"]["scan_complete"] is True
    assert result["state"] == "PARTIAL"  # a selective materialization is never reported as complete
    assert any("selective_materialization" in r for r in result["reasons"])
    row = document.passages[-1]
    assert row.locator == "delimited:row=5002" and row.text == "Aberrant_Basaloid | TEAD1 | 0.42 | 0.61"
    assert [c["header"] for c in row.cells] == ["cellType", "gene", "log1p_FC", "pct.inGroup"]
    # bounded by construction: passages never exceed the sample plus the selector cap, whatever the table size
    assert len(document.passages) <= 2000 + 1 + 1


def test_selector_matches_beyond_the_cap_are_counted_not_silently_dropped():
    from regkg.literature.files import RowSelector, parse_delimited

    content = ("gene\tv\n" + "".join(f"TEAD1\t{i}\n" for i in range(50))).encode()
    selector = RowSelector("test", frozenset({"TEAD1"}), frozenset())
    _, report = parse_delimited("pub:1", content, "t.tsv", "\t", 2, 10, selector, max_selected_rows=5)
    result = report.to_json()
    assert result["components"]["rows_matching_selector"] == 50
    assert result["components"]["selector_rows_materialized_beyond_sample"] == 5
    assert any("selector_match_limit_reached" in r for r in result["reasons"])


def test_a_named_size_allowance_reopens_only_its_own_supplement():
    from regkg.workflows.corpus_prepare import _dependency_state

    deferred = {"state": "size_limit_deferred", "filename": "big.zip", "size": 32_000_000}
    assert _dependency_state(deferred, []) == "deferred_size_limit"
    assert _dependency_state(deferred, [], (), 67_108_864) == "to_acquire"
    assert _dependency_state({**deferred, "size": 90_000_000}, [], (), 67_108_864) == "deferred_size_limit"


def test_archive_expansion_stays_bounded_and_unsafe_members_are_still_refused(tmp_path):
    import io
    import zipfile

    from regkg.literature.files import inspect_zip

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("table_s4.csv", "gene,value\nKLF5,1\n")
        archive.writestr("../escape.csv", "gene,value\nX,1\n")
        archive.writestr("big.csv", "x" * 5000)
    members, report = inspect_zip(buffer.getvalue(), "supplement.zip", max_bytes=1000)
    result = report.to_json()
    dispositions = {m["name"]: m["disposition"] for m in result["members"]}
    assert dispositions["table_s4.csv"] == "extracted_for_routing"
    assert dispositions["../escape.csv"] == "rejected_path_traversal"
    assert dispositions["big.csv"] == "deferred_archive_limit"  # the allowance bounds expansion
    assert [name for name, _ in members] == ["table_s4.csv"] and result["state"] == "PARTIAL"


def test_parse_cache_identity_includes_the_selection_and_the_effective_limits(tmp_path):
    """The defect the lead reproduced: selector A then selector B must not return A's cached rows."""
    from regkg.literature.files import RowSelector, parse_delimited
    from regkg.workflows.corpus_parse import delimited_options, parse_cached

    content = b"gene,value\nTEAD1,1\nKLF5,2\nRFX2,3\n"
    calls = []

    def run(selector, max_rows=1):
        own = SimpleNamespace(max_rows_per_sheet=max_rows, max_columns=10, max_sheets=5)

        def parse():
            calls.append(selector.rule if selector else None)
            return parse_delimited("pub:1", content, "t.csv", ",", max_rows, 10, selector, 2000)

        document, _ = parse_cached(
            tmp_path,
            "delimited",
            "pub:1",
            content,
            "t.csv",
            lambda: _json(parse()),
            delimited_options(own, selector, 2000),
        )
        return document

    def _json(result):
        document, report = result
        return document, report.to_json()

    a = RowSelector("sel", frozenset({"TEAD1"}), frozenset())
    b = RowSelector("sel", frozenset({"RFX2"}), frozenset())  # same counts, different contents
    assert a.fingerprint()["digest"] != b.fingerprint()["digest"]
    first = run(a)
    second = run(b)
    assert len(calls) == 2, "selector B must not reuse selector A's cache entry"
    assert "TEAD1" in first.passages[-1].text and "RFX2" in second.passages[-1].text
    run(b)
    assert len(calls) == 2  # identical settings reuse the cache
    run(b, max_rows=3)
    assert len(calls) == 3  # changed row limit invalidates
    # selection identity reaches the parsed document's own rules, so its document id differs too
    assert first.extra_rules["selector"]["digest"] != second.extra_rules["selector"]["digest"]
    assert first.extra_rules["max_rows"] == 1
